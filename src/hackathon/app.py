import datetime
import io
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

load_dotenv()

EXHIBIT3_PATH = Path("synthetic_data/exhibit3.xlsx")
MODEL = "gpt-5-mini"
TEST_SHEET = re.compile(r"Module\s*(\d+)\s*Test\s*([A-C])", re.IGNORECASE)
GRID_COLS = 18

# status -> (xlsx/css hex fill, text hex, legend label)
STATUS_COLORS = {
    "mastered": ("63BE7B", "000000", "Completely mastered"),
    "partial": ("FFD666", "000000", "Partially mastered"),
    "not_understood": ("F8696B", "000000", "Not understood"),
    "remaining": ("595959", "FFFFFF", "Remaining (not yet covered)"),
}


def score_results(results):
    """Score a test: DONE / total questions, C and blank count as 0."""
    if not results:
        return None
    done = sum(1 for r in results if r.strip().upper() == "DONE")
    return round(100 * done / len(results))


def parse_student_id(filename):
    m = re.match(r"(\d+)_exhibit1", filename, re.IGNORECASE)
    return m.group(1) if m else None


def parse_exhibit1(file):
    """Return (scores, points). scores: {(module, test): int}. points: list of
    (module, test, syllabus_point, result)."""
    wb = load_workbook(file, data_only=True, read_only=True)
    buckets = defaultdict(list)
    points = []
    for ws in wb.worksheets:
        m = TEST_SHEET.match(ws.title)
        if not m:
            continue
        module, test = int(m.group(1)), m.group(2).upper()
        header = None
        for row in ws.iter_rows(min_row=1, max_row=10):
            labels = [str(c.value).strip().lower() if c.value else "" for c in row]
            if "result" in labels:
                header = row[0].row
                key_col = labels.index("key") + 1 if "key" in labels else 1
                point_col = (
                    labels.index("syllabus point") + 1
                    if "syllabus point" in labels
                    else key_col
                )
                result_col = labels.index("result") + 1
                break
        if header is None:
            continue
        for row in ws.iter_rows(min_row=header + 1):
            key = ws.cell(row[0].row, key_col).value
            if key is None:
                continue
            result = ws.cell(row[0].row, result_col).value
            result = "" if result is None else str(result).strip()
            buckets[(module, test)].append(result)
            point = ws.cell(row[0].row, point_col).value or key
            points.append((module, test, str(point), result))
    scores = {k: score_results(v) for k, v in buckets.items()}
    return scores, points


def exhibit3_colmap(ws):
    """Map (module, test) -> column index, read from the module/test headers."""
    module_at = {}
    current = None
    for cell in ws[5]:
        m = re.match(r"Module\s*(\d+)", str(cell.value or ""))
        if m:
            current = int(m.group(1))
        module_at[cell.column] = current
    colmap = {}
    for cell in ws[6]:
        letter = str(cell.value or "").strip().upper()
        if letter in ("A", "B", "C") and module_at.get(cell.column):
            colmap[(module_at[cell.column], letter)] = cell.column
    return colmap


def update_exhibit3(student_id, scores):
    """Write the student's scores into exhibit3 in place."""
    wb = load_workbook(EXHIBIT3_PATH)
    ws = wb["Exhibit 3"] if "Exhibit 3" in wb.sheetnames else wb.active
    colmap = exhibit3_colmap(ws)
    student_row = None
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value is not None and str(row[0].value).strip() == str(student_id):
            student_row = row[0].row
            break
    if student_row is None:
        raise ValueError(f"Student {student_id} not found in exhibit3.")
    for (module, test), score in scores.items():
        col = colmap.get((module, test))
        if col and score is not None:
            ws.cell(student_row, col).value = score
    wb.save(EXHIBIT3_PATH)


def _point_key(label):
    """Natural sort key for dotted syllabus points like 1.2.10."""
    return [int(p) if p.isdigit() else p for p in re.split(r"[.\-]", label)]


def aggregate_points(points):
    """Collapse per-test results into one status per syllabus point.

    DONE -> mastered, C -> partial, blank -> remaining, anything else ->
    not understood, aggregated across a point's A/B/C tests."""
    by_point = defaultdict(list)
    for _module, _test, point, result in points:
        by_point[point].append(result)
    aggregated = []
    for point in sorted(by_point, key=_point_key):
        attempts = [r.strip().upper() for r in by_point[point] if r.strip()]
        if not attempts:
            status = "remaining"
        elif all(r == "DONE" for r in attempts):
            status = "mastered"
        elif any(r in ("DONE", "C") for r in attempts):
            status = "partial"
        else:
            status = "not_understood"
        aggregated.append((point, status))
    return aggregated


def point_summary(aggregated):
    total = len(aggregated)
    counts = Counter(status for _point, status in aggregated)
    covered = total - counts["remaining"]
    understood = counts["mastered"] + counts["partial"]
    mastered = counts["mastered"]

    def pct(numerator, denominator):
        return round(100 * numerator / denominator, 2) if denominator else 0.0

    return {
        "total": total,
        "covered": covered,
        "covered_pct": pct(covered, total),
        "understood": understood,
        "understood_pct": pct(understood, total),
        "mastered": mastered,
        "mastered_pct": pct(mastered, total),
        "mastered_covered_pct": pct(mastered, covered),
    }


def heatmap_html(student_id, aggregated):
    cells = "".join(
        f'<div style="background:#{STATUS_COLORS[status][0]};'
        f"color:#{STATUS_COLORS[status][1]};padding:6px 4px;width:60px;"
        f'text-align:center;font-size:12px;border-radius:2px">{point}</div>'
        for point, status in aggregated
    )
    legend = "".join(
        f'<div style="display:flex;align-items:center;gap:6px">'
        f'<span style="width:14px;height:14px;background:#{bg};border-radius:2px"></span>'
        f'<span style="font-size:12px">{label}</span></div>'
        for bg, _fg, label in STATUS_COLORS.values()
    )
    return (
        f'<div style="font-weight:600;margin-bottom:6px">Student {student_id} '
        f"— MyTalent progress report</div>"
        f'<div style="display:flex;gap:16px;flex-wrap:wrap;margin-bottom:12px">{legend}</div>'
        f'<div style="display:flex;flex-wrap:wrap;gap:3px">{cells}</div>'
    )


def build_exhibit2(student_id, aggregated, summary):
    """Grid heatmap workbook matching the progress-report layout. Returns bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Progress report"
    fills = {k: PatternFill("solid", fgColor=v[0]) for k, v in STATUS_COLORS.items()}
    ws["A1"] = f"Student {student_id} — MyTalent progress report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Last updated {datetime.date.today():%Y/%m/%d}"
    for i, (status, meta) in enumerate(STATUS_COLORS.items()):
        cell = ws.cell(3, 1 + i, meta[2])
        cell.fill = fills[status]
    start = 5
    for idx, (point, status) in enumerate(aggregated):
        row, col = divmod(idx, GRID_COLS)
        cell = ws.cell(start + row, 1 + col, point)
        cell.fill = fills[status]
        cell.alignment = Alignment(horizontal="center")
    summary_row = start + (len(aggregated) // GRID_COLS) + 3
    rows = [
        ("Total points", summary["total"], ""),
        ("Covered points", summary["covered"], f"{summary['covered_pct']}%"),
        ("Content understood", summary["understood"], f"{summary['understood_pct']}%"),
        (
            "Content mastered (all points)",
            summary["mastered"],
            f"{summary['mastered_pct']}%",
        ),
        (
            "Content mastered (covered points)",
            "",
            f"{summary['mastered_covered_pct']}%",
        ),
    ]
    for i, (label, value, percent) in enumerate(rows):
        ws.cell(summary_row + i, 1, label)
        ws.cell(summary_row + i, 2, value)
        ws.cell(summary_row + i, 3, percent)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def exhibit2_context(student_id, summary, aggregated):
    lines = [
        f"Student {student_id} progress report (exhibit2):",
        (f"Total points {summary['total']}, covered {summary['covered']} "
        f"({summary['covered_pct']}%), understood {summary['understood']} "
        f"({summary['understood_pct']}%), mastered {summary['mastered']} "
        f"({summary['mastered_pct']}%)."),
    ]
    for status in ("partial", "not_understood", "remaining"):
        pts = [point for point, s in aggregated if s == status]
        if pts:
            lines.append(f"{STATUS_COLORS[status][2]}: " + ", ".join(pts))
    return "\n".join(lines)


@st.cache_resource
def openai_client():
    return OpenAI()


st.title("Hi, welcome to MyTalent")

uploaded = st.file_uploader(
    "Upload a student exhibit1 file",
    type=["xlsx"],
    help="Filename format: {student_id}_exhibit1.xlsx",
)

if uploaded is not None:
    file_id = (uploaded.name, uploaded.size)
    if st.session_state.get("file_id") != file_id:
        student_id = parse_student_id(uploaded.name)
        if student_id is None:
            st.error("Filename must look like {student_id}_exhibit1.xlsx")
            st.stop()
        try:
            scores, points = parse_exhibit1(uploaded)
            update_exhibit3(student_id, scores)
        except Exception as exc:  # surface processing errors to the tutor
            st.error(f"Could not process this file: {exc}")
            st.stop()
        aggregated = aggregate_points(points)
        summary = point_summary(aggregated)
        st.session_state.file_id = file_id
        st.session_state.student_id = student_id
        st.session_state.aggregated = aggregated
        st.session_state.summary = summary
        st.session_state.exhibit2 = build_exhibit2(student_id, aggregated, summary)
        st.session_state.context = exhibit2_context(student_id, summary, aggregated)
        st.session_state.messages = []
        st.toast("exhibit3.xlsx has been updated", icon=":material/check_circle:")

if "student_id" in st.session_state:
    sid = st.session_state.student_id
    st.success(f"exhibit3.xlsx has been updated for student {sid}.")

    st.subheader("Exhibit 2 — progress report")
    summary = st.session_state.summary
    st.caption(f"Total points in this upload: {summary['total']}")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric(
        "Covered points",
        summary["covered"],
        f"{summary['covered_pct']}%",
        delta_color="off",
    )
    c2.metric(
        "Content understood",
        summary["understood"],
        f"{summary['understood_pct']}%",
        delta_color="off",
    )
    c3.metric(
        "Content mastered",
        summary["mastered"],
        f"{summary['mastered_pct']}%",
        delta_color="off",
    )
    c4.metric("Mastered of covered", f"{summary['mastered_covered_pct']}%")

    st.html(heatmap_html(sid, st.session_state.aggregated))

    st.download_button(
        "Download exhibit2 heatmap",
        data=st.session_state.exhibit2,
        file_name=f"{sid}_exhibit2.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    st.subheader("Ask about this student's plan")
    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    if prompt := st.chat_input(
        "Ask a follow-up to build a personalised plan", submit_mode="disable"
    ):
        if not os.environ.get("OPENAI_API_KEY"):
            st.error("OPENAI_API_KEY is not set in .env")
            st.stop()
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)
        system = (
            "You are a MyTalent tutor assistant. Use the student's exhibit2 results "
            "below to answer the tutor's questions and build a concrete, personalised "
            "study plan that targets the weak syllabus points.\n\n"
            + st.session_state.context
        )
        with st.chat_message("assistant"):
            stream = openai_client().chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": system},
                    *st.session_state.messages,
                ],
                stream=True,
            )
            reply = st.write_stream(stream)
        st.session_state.messages.append({"role": "assistant", "content": reply})


if __name__ == "__main__" and os.environ.get("RUN_SELFTEST"):
    assert score_results(["DONE", "DONE", "C", "", "W"]) == 40
    assert score_results(["DONE", "DONE"]) == 100
    assert score_results([]) is None
    assert parse_student_id("3_exhibit1.xlsx") == "3"
    assert parse_student_id("bad.xlsx") is None
    _points = [
        (1, "A", "1.1", "DONE"),
        (1, "B", "1.1", "DONE"),  # mastered
        (1, "A", "1.2", "C"),  # partial
        (1, "A", "1.10", "W"),  # not understood
        (1, "A", "1.3", ""),  # remaining
    ]
    _agg = dict(aggregate_points(_points))
    assert _agg == {
        "1.1": "mastered",
        "1.2": "partial",
        "1.10": "not_understood",
        "1.3": "remaining",
    }
    assert [p for p, _ in aggregate_points(_points)] == [
        "1.1",
        "1.2",
        "1.3",
        "1.10",
    ]  # natural sort
    _s = point_summary(aggregate_points(_points))
    assert (_s["total"], _s["covered"], _s["understood"], _s["mastered"]) == (
        4,
        3,
        2,
        1,
    )
    print("selftest ok")
